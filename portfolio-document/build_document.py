from html.parser import HTMLParser
from pathlib import Path
import re,json,html
from urllib.parse import urljoin
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
class N:
 def __init__(self,tag='',attrs=[]): self.tag=tag;self.a=dict(attrs);self.children=[]
 def text(self): return re.sub(r'\s+',' ',' '.join(c if isinstance(c,str) else c.text() for c in self.children)).strip()
 def find(self,cls=None,tag=None):
  out=[]
  for c in self.children:
   if isinstance(c,N):
    if (cls is None or cls in c.a.get('class','').split()) and (tag is None or c.tag==tag):out.append(c)
    out+=c.find(cls,tag)
  return out
class Parser(HTMLParser):
 def __init__(self):super().__init__();self.root=N();self.stack=[self.root]
 def handle_starttag(self,t,a):
  n=N(t,a);self.stack[-1].children.append(n)
  if t not in ['img','br','hr','meta','link','input','source','wbr']:self.stack.append(n)
 def handle_endtag(self,t):
  for i in range(len(self.stack)-1,0,-1):
   if self.stack[i].tag==t:self.stack=self.stack[:i];break
 def handle_data(self,d):self.stack[-1].children.append(d)
p=Parser();p.feed(Path('portfolio-document/source.html').read_text(encoding='utf-8'));root=p.root
get=lambda cls:root.find(cls)[0].text()
def plain(s):return re.sub(r'\s+',' ',html.unescape(re.sub('<[^>]+>',' ',s))).strip()
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
sec.page_height=Inches(11.7);sec.page_width=Inches(8.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(25)
for name,size in [('Heading 1',17),('Heading 2',12),('Heading 3',11)]:doc.styles[name].font.size=Pt(size)
def para(t,style=None): return doc.add_paragraph(t,style)
def heading(t,l=1):doc.add_heading(re.sub(r'[^\w\s]',' ',t).replace('_',' '),l)
def label(a,b):
 p=para('');p.add_run(a+': ').bold=True;p.add_run(b)
def link(label_,url):
 p=para('');p.add_run(label_+': ').bold=True
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True));r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=url;r.append(t);h.append(r);p._p.append(h)
para('Devasish Portfolio Details','Title')
para('Devasish Venkat Sai Jajimoggala','Subtitle')
para('A complete record of the Devaverse portfolio, including the profile, education, seven projects, detailed project showcases, skills, certifications, contact information, and website links.')
link('Portfolio','https://devaverse.netlify.app');label('Captured','9 October 2026')
heading('Profile')
label('Website headline',get('hero-name'));label('Professional focus',get('hero-eyebrow'));label('Education summary',get('hero-role'));para(get('hero-desc'));para(get('about-body'))
for row in root.find('about-row'):label(row.find('lbl')[0].text(),row.find('val')[0].text())
label('Portfolio statistics','7 projects built | 10+ certifications | 12+ tech stacks')
label('Areas highlighted','Full Stack Dev; Cyber Security; Flutter & Mobile; Node.js & Express; PostgreSQL & MongoDB; React & UI Design; JWT & WebSockets')
heading('Education')
para(get('edu-deg'));para(get('edu-uni'));para(get('edu-year'));label('Relevant coursework','; '.join(x.text() for x in root.find('ctag')))
heading('Security philosophy')
para(get('rs-title'));para(get('rs-text'))
projects=json.loads(Path('portfolio-document/projects.json').read_text(encoding='utf-8'))
for card in root.find('proj-card'):
 key=re.search("openModal\('([^']+)'",card.a['onclick']).group(1);data=projects[key]
 doc.add_page_break();heading(card.find('proj-name')[0].text())
 label('Portfolio listing',card.find('proj-num')[0].text()+' | '+card.find('proj-year')[0].text())
 if card.find('proj-highlight-badge'):label('Status','Featured')
 para(card.find('proj-desc')[0].text())
 heading('Project card highlights',2)
 for li in card.find(tag='li'):para(li.text(),'List Bullet')
 label('Card technology stack','; '.join(x.text() for x in card.find('stack-tag')))
 heading('Detailed project description',2)
 label('Showcase category',data['cat']);label('Showcase number',data['num'])
 for s in data['about']:para(plain(s))
 heading('Detailed features',2)
 for f in data['features']:para(f,'List Bullet')
 label('Full technology stack','; '.join(data['stack']))
 for k,title in [('github','GitHub repository'),('demo','Live page')]:
  if k in data:link(title,data[k])
 heading('Website showcase preview',2)
 para('The following text is displayed in the portfolio preview, including illustrative status and counter values.')
 label('Preview title',data['previewTitle'])
 for line in re.split(r'<br\s*/?>',data['preview']):
  txt=plain(line)
  txt=re.sub('[\U00010000-\U0010ffff\ufe0f]','',txt).replace('✓','[OK]').replace('●','').replace('⚡','').replace('⚠','[!]')
  if txt.strip():para(txt.strip())
 for s in data.get('screenshots',[]):link('Showcase screenshot',urljoin('https://devaverse.netlify.app/',s))
doc.add_page_break();heading('Skills and expertise')
para(root.find('skills-intro')[0].find(tag='p')[0].text())
for group in root.find('skills-group'):
 heading(group.find('skills-group-title')[0].text(),2);para('; '.join(x.text() for x in group.find(tag='li')))
heading('Certifications')
for cert in root.find('cert-card'):label(cert.find('cert-org')[0].text(),cert.find('cert-name')[0].text())
heading('Contact and availability')
contact=[n for n in root.find(tag='section') if n.a.get('id')=='contact'][0]
for x in contact.find(tag='h2')+contact.find(tag='p'):para(x.text())
link('Email','mailto:jajimoggaladevasishvenkatsai@gmail.com');label('Phone','+91 9505271744');link('GitHub','https://github.com/Devasish009');link('LinkedIn','https://linkedin.com/in/devasishjajimoggala');link('Devaverse','https://devaverse.netlify.app');link('Download resume','https://devaverse.netlify.app/assets/Devasish_Resume.pdf')
heading('Website navigation and actions')
para('Navigation: About, Projects, Skills, Certs, Contact. Homepage actions: Hire Me, View Projects, Download Resume, Get in Touch, Email Me, GitHub, LinkedIn, Devaverse, and Lets build Something. Project showcases provide View on GitHub, Live page where available, and Close actions. The theme control switches between Cyber Mode and Crimson Mode. The page includes a Scroll prompt and the brand DEVASISH.')
label('Section labels','About Me / Who I Am; Work / Featured Projects; Expertise / My Tech Arsenal; Credentials / Certifications; Let’s Connect / Got an Idea? Let’s Build It.')
para('Featured projects tagline: Real products. Real tech. Real impact.')
para('Footer: © 2026 Devasish Venkat Sai Jajimoggala. Footer links: GitHub, LinkedIn, devaverse.')
out=Path('portfolio-document/Devasish_Portfolio_Details.docx').resolve();doc.save(out);print(out)
print('Projects:',len(root.find('proj-card')),'Skill groups:',len(root.find('skills-group')),'Certificate entries:',len(root.find('cert-card')))
